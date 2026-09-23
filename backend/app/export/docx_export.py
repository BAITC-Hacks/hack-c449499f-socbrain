import io

from docx import Document
from docx.shared import Pt

from ..config import settings
from . import speaker_names, summary_of, task_rows, timecode


def build(meeting: dict) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)

    doc.add_heading("Протокол совещания", level=0)
    org = settings.values["appearance"]["org_name"]
    if org:
        doc.add_paragraph().add_run(org).bold = True
    doc.add_paragraph(f"Тема: {meeting['title']}\nДата: {meeting['meeting_date']}")

    names = speaker_names(meeting)
    doc.add_heading("Участники", level=1)
    for s in meeting["speakers"]:
        doc.add_paragraph(f"{s['name'] or s['label']}" + (f" — {s['role']}" if s["role"] else ""),
                          style="List Bullet")

    summary = summary_of(meeting)
    doc.add_heading("Краткое содержание", level=1)
    if summary.get("overview"):
        doc.add_paragraph(summary["overview"])
    for topic in summary.get("topics", []):
        doc.add_paragraph().add_run(topic["title"]).bold = True
        for point in topic.get("points", []):
            doc.add_paragraph(point, style="List Bullet")
    if summary.get("decisions"):
        doc.add_paragraph().add_run("Решения").bold = True
        for decision in summary["decisions"]:
            doc.add_paragraph(decision, style="List Bullet")

    doc.add_heading("Поручения", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, ["№", "Поручение", "Ответственный", "Срок"]):
        cell.text = title
        cell.paragraphs[0].runs[0].bold = True
    for values in task_rows(meeting):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value

    doc.add_heading("Стенограмма", level=1)
    for seg in meeting["segments"]:
        p = doc.add_paragraph()
        p.add_run(f"[{timecode(seg['start'])}] {names.get(seg['speaker'], seg['speaker'])}: ").bold = True
        p.add_run(seg["text"])

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
